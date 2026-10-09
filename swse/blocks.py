"""Block specifications: where each table lives inside a source workbook.

A *block* is a rectangular region of a worksheet that holds one logical table.
Both source workbooks pack many blocks per sheet (SagaForge's hidden ``Data``
sheet alone holds ~30), so blocks are described declaratively in
``config/blocks.yaml`` and read by :mod:`swse.extract`.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

import yaml
from openpyxl.utils import column_index_from_string, get_column_letter

from . import paths
from .ids import clean_cell, fold


@dataclass
class Block:
    id: str
    entity: str
    source: str
    sheet: str
    #: Row that holds column headers (may be partially empty).
    header_row: int = 1
    #: First row of data.
    first_row: int = 2
    #: Explicit last data row; ``None`` -> auto-detect.
    last_row: int | None = None
    #: Column letters in scope, e.g. ``["C", "D", "E"]`` (expanded from ``"C:E"``).
    columns: list[str] = field(default_factory=list)
    #: Column letter -> field name. Wins over the header text.
    rename: dict[str, str] = field(default_factory=dict)
    #: Header text (lowercased) -> field name. Used for messy headers.
    rename_header: dict[str, str] = field(default_factory=dict)
    #: Columns whose value should be carried down through blank rows.
    forward_fill: list[str] = field(default_factory=list)
    #: Column whose blankness marks the end of the table during auto-detect.
    key_column: str | None = None
    #: When true, rows with a blank `key_column` are skipped entirely. Use for
    #: blocks that share a row range with unrelated helper columns.
    require_key: bool = False
    #: Stop auto-detect after this many consecutive empty rows.
    blank_run: int = 8
    #: Field(s) used to synthesise a display `name` when the block has none.
    name_field: list[str] = field(default_factory=list)
    #: Optional ``"Label {value}"`` template applied to the synthesised name.
    name_template: str = ""
    canon: str = "official"
    notes: str = ""
    enabled: bool = True

    # ------------------------------------------------------------------ helpers
    def field_for(self, col: str, header: Any) -> str:
        """Field name for a column: explicit rename > header text > letter."""
        if col in self.rename:
            return self.rename[col]
        h = fold(header).lower()
        if h in self.rename_header:
            return self.rename_header[h]
        if h:
            slug = "".join(ch if ch.isalnum() else "_" for ch in h)
            while "__" in slug:
                slug = slug.replace("__", "_")
            return slug.strip("_") or f"col_{col.lower()}"
        return f"col_{col.lower()}"

    def ref(self, row: int) -> str:
        if not self.columns:
            return f"{self.sheet}!{row}"
        return f"{self.sheet}!{self.columns[0]}{row}:{self.columns[-1]}{row}"


def _expand_columns(spec: Any) -> list[str]:
    """``"C:N"`` -> ``["C", ..., "N"]``; ``["C","D"]`` -> unchanged."""
    if spec is None:
        return []
    if isinstance(spec, str):
        if ":" in spec:
            a, b = spec.split(":", 1)
            return [get_column_letter(i) for i in range(column_index_from_string(a), column_index_from_string(b) + 1)]
        return [spec]
    out: list[str] = []
    for item in spec:
        out.extend(_expand_columns(item))
    return out


def load_blocks(only_enabled: bool = True) -> list[Block]:
    doc = yaml.safe_load(paths.BLOCKS_YAML.read_text(encoding="utf-8")) or {}
    blocks = []
    for entry in doc.get("blocks", []):
        b = Block(
            id=entry["id"],
            entity=entry.get("entity", entry["id"]),
            source=entry["source"],
            sheet=entry["sheet"],
            header_row=int(entry.get("header_row", 1)),
            first_row=int(entry.get("first_row", 2)),
            last_row=int(entry["last_row"]) if entry.get("last_row") else None,
            columns=_expand_columns(entry.get("columns")),
            rename=dict(entry.get("rename") or {}),
            rename_header={str(k).lower(): v for k, v in (entry.get("rename_header") or {}).items()},
            forward_fill=list(entry.get("forward_fill") or []),
            key_column=entry.get("key_column"),
            require_key=bool(entry.get("require_key", False)),
            blank_run=int(entry.get("blank_run", 8)),
            name_field=(entry.get("name_field") if isinstance(entry.get("name_field"), list)
                        else ([entry["name_field"]] if entry.get("name_field") else [])),
            name_template=entry.get("name_template", ""),
            canon=entry.get("canon", "official"),
            notes=entry.get("notes", ""),
            enabled=bool(entry.get("enabled", True)),
        )
        if only_enabled and not b.enabled:
            continue
        blocks.append(b)
    ids = [b.id for b in blocks]
    dupes = {i for i in ids if ids.count(i) > 1}
    if dupes:
        raise ValueError(f"duplicate block ids in blocks.yaml: {sorted(dupes)}")
    return blocks


def get_block(block_id: str) -> Block:
    for b in load_blocks():
        if b.id == block_id:
            return b
    raise KeyError(f"unknown block {block_id!r}")


def cell_value(ws, row: int, col: str):
    return clean_cell(ws[f"{col}{row}"].value)
